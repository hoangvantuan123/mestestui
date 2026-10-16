from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path(__file__).parent
d = Document()
s = d.sections[0]
s.page_height, s.page_width = Cm(29.7), Cm(21)
s.top_margin = s.bottom_margin = Cm(1.8)
s.left_margin = s.right_margin = Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=d.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.12
d.styles['Title'].font.size=Pt(24)
d.styles['Heading 1'].font.size=Pt(17)
d.styles['Heading 2'].font.size=Pt(12)
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def table(headers, rows, widths):
    t=d.add_table(rows=1, cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for c,txt in zip(t.rows[0].cells,headers): c.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row): c.text=txt
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j]); c.vertical_alignment=1
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            pr.append(borders)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE6F1' if i==0 else ('F5F7FA' if i%2==0 else 'FFFFFF')); pr.append(shade)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'100'); el.set(qn('w:type'),'dxa'); mar.append(el)
            pr.append(mar)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                para.paragraph_format.line_spacing=1.05
                for r in para.runs: r.font.size=Pt(10); r.bold=(i==0)
    rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    d.add_paragraph().paragraph_format.space_after=Pt(0)

d.add_paragraph('Quy trình xử lý Rework và Sorting trên MES và BRAVO','Title')
p('Mô tả luồng nghiệp vụ và các điểm kiểm tra chứng từ')
d.add_heading('1 Phạm vi nghiệp vụ',1)
p('Tài liệu mô tả điểm bắt đầu, các bước xử lý trên MES và BRAVO, luồng điều chuyển vật tư và kết quả chứng từ của nghiệp vụ Rework. Người vận hành MES, bộ phận kho và người kiểm tra BRAVO sử dụng tài liệu để đối chiếu lệnh, mã vật tư, kho và tài khoản trước khi hoàn tất xử lý.')
table(['Nhóm','Mã','Nghiệp vụ'],[
('Rework','8A','Rework nội bộ'),('Rework','8B','Rework hàng bán bị trả lại để bán lại theo RMA'),
('Sorting','7A','Sorting nội bộ'),('Sorting','7B','Sorting hàng bán bị trả lại theo RMA')],[3,1.5,12.5])
p('Luồng chi tiết dưới đây áp dụng chung cho 8A và 8B. Hai mã phân biệt nguồn nghiệp vụ nội bộ và hàng trả lại RMA. Các yêu cầu riêng về tiếp nhận hoặc liên kết RMA được đưa vào mục thông tin cần xác nhận.')
h('Luồng tổng quát của 8A và 8B')
p('BRAVO tạo lệnh công đoạn → MES đồng bộ lệnh về MES Admin và MES xưởng → Máy đại diện Rework thực hiện Lấy NVL → MES đẩy phiếu đề nghị xuất kho lên BRAVO → BRAVO điều chuyển từ kho chính sang kho phân xưởng → Thực hiện Rework → MES thống kê và hoàn thành lệnh → Dữ liệu hoàn thành chuyển lên BRAVO → BRAVO tự động sinh phiếu xuất chuyển mã.')
p('Nguyên tắc kiểm soát: phiếu đề nghị xuất kho chưa đồng nghĩa với vật tư đã được điều chuyển. Trước khi hoàn thành lệnh, phải bảo đảm mã hàng cần xử lý còn tồn tại kho phân xưởng.')

page('2 Tạo lệnh và lấy nguyên vật liệu')
h('Bước 1 Tạo lệnh công đoạn trên BRAVO')
p('Tạo lệnh công đoạn thuộc nghiệp vụ 8A hoặc 8B trên BRAVO. Mã vật tư chính đăng ký tại lệnh là căn cứ xác định vật tư Rework được đề nghị xuất và vật tư xuất khi chuyển mã.')
h('Bước 2 Đồng bộ lệnh sang MES')
p('MES đồng bộ lệnh từ BRAVO về MES Admin và MES xưởng. Kiểm tra lệnh đã xuất hiện đúng trên MES trước khi thao tác. Người vận hành truy cập máy đại diện Rework để thực hiện Lấy NVL và Thống kê.')
h('Bước 3 Lấy NVL trên MES')
p('Tại máy đại diện Rework, chọn đúng lệnh và thực hiện Lấy NVL. MES đẩy phiếu đề nghị xuất kho lên BRAVO. Kiểm tra phiếu đã được tạo thành công và khớp với lệnh đang xử lý.')
h('Thông tin đầu phiếu đề nghị xuất kho trên BRAVO')
table(['Mục','Thông tin cần kiểm tra'],[
('4 Đối tượng','Kiểm tra đối tượng gắn với nghiệp vụ và lệnh đang thực hiện; giá trị cụ thể cần xác nhận theo cấu hình thực tế.'),
('6 Công đoạn','Chọn đúng CDRE thuộc nhóm Rework.'),
('7 Loại xuất','X01 — Xuất theo lệnh.'),
('9 Kho nhập','Tự động xác định ra kho phân xưởng. Kiểm tra đúng kho nhận trước khi điều chuyển.')],[4,13])
h('Thông tin chi tiết đề nghị xuất kho')
p('Đề nghị xuất sử dụng chính mã vật tư Rework của lệnh, không lấy danh sách NVL trong BOM làm căn cứ cho dòng đề nghị xuất này. Kiểm tra mã vật tư, số lượng đề nghị và kho xuất; đặc biệt phải kiểm tra tồn của đúng mã vật tư tại kho xuất.')
p('Điểm kiểm tra trước khi chuyển bước: đúng lệnh, đúng mã Rework, đúng CDRE, đúng X01, đúng kho xuất và kho phân xưởng nhận; tồn kho xuất đáp ứng lượng cần điều chuyển.')

page('3 Điều chuyển kho và hoàn thành lệnh')
h('Bước 4 Điều chuyển trên BRAVO')
p('Sau khi có phiếu đề nghị xuất kho, phía BRAVO sử dụng phiếu này để thực hiện điều chuyển vật tư từ kho chính sang kho phân xưởng phục vụ Rework. Đối chiếu mã vật tư, số lượng và hai kho theo đề nghị xuất.')
p('Kết quả cần đạt: vật tư đã được chuyển sang kho phân xưởng và tồn kho phân xưởng đã phản ánh đủ lượng cần xử lý. Việc MES lấy NVL thành công và BRAVO có phiếu đề nghị xuất là mốc tạo yêu cầu; cần kiểm tra kết quả điều chuyển thực tế trước khi hoàn thành lệnh.')
h('Bước 5 Thực hiện Rework tại phân xưởng')
p('Phân xưởng thực hiện Rework theo lệnh. Sau khi xử lý, người vận hành sử dụng chức năng Thống kê và Hoàn thành lệnh trên MES tại máy đại diện Rework.')
h('Bước 6 Kiểm tra tồn trước khi hoàn thành trên MES')
p('Trước khi bấm Hoàn thành lệnh, kiểm tra tồn của đúng mã hàng tại đúng kho phân xưởng. Không tiếp tục hoàn thành khi chưa bảo đảm tồn kho đáp ứng số lượng cần xử lý; cần đối chiếu lại phiếu điều chuyển, mã hàng và mã kho.')
p('Các ký hiệu trường “b. Mã hàng” và “d. Mã kho phân xưởng” cần được đối chiếu trên màn hình thực tế để xác nhận cách ghép mã hàng và kho khi kiểm tra tồn. Không suy diễn “mã hàng kho chính” thành một mã vật tư khác khi chưa xác nhận ánh xạ.')
h('Bước 7 Chuyển kết quả hoàn thành sang BRAVO')
p('Khi Thống kê và Hoàn thành lệnh trên MES thành công, dữ liệu hoặc phiếu hoàn thành được đẩy lên BRAVO. BRAVO tự động sinh phiếu xuất chuyển mã theo kho xuất và kho nhập đã chỉ định trước đó.')
h('Thông tin đầu phiếu xuất chuyển mã')
table(['Trường','Nội dung đối chiếu'],[
('K Kho xuất','Đúng kho xuất được chỉ định cho giao dịch chuyển mã.'),
('J Kho nhập','Đúng kho nhập được chỉ định cho giao dịch chuyển mã.'),
('9 Giao dịch','2212 — Xuất chuyển mã cho Rework tự động.')],[4,13])

page('4 Kiểm tra chuyển mã và nội dung cần xác nhận')
h('Chi tiết phiếu xuất chuyển mã trên BRAVO')
table(['Nhóm chi tiết','Nội dung cần kiểm tra'],[
('Vật tư xuất','Mã vật tư xuất là mã vật tư chính đã đăng ký trên lệnh công đoạn. Đối chiếu mã, số lượng, TK nợ và TK có.'),
('Vật tư nhập','Mã vật tư nhập là danh sách NVL chính cấu thành nên hàng đã khai báo ở Rework. Đối chiếu từng mã với khai báo Rework; kiểm tra số lượng và các tài khoản liên quan theo cấu hình.')],[4,13])
p('Cần phân biệt hai thời điểm: khi Lấy NVL, đề nghị xuất dùng chính mã Rework của lệnh; khi xuất chuyển mã, nhóm vật tư nhập là danh sách NVL chính được khai báo cấu thành hàng Rework. Không tự thay danh sách vật tư nhập bằng BOM khi chưa đối chiếu quy tắc khai báo Rework.')
h('Danh sách kiểm tra kết thúc luồng 8A và 8B')
for t in [
'1. Lệnh BRAVO đã đồng bộ đúng trên MES Admin và MES xưởng.',
'2. Lấy NVL trên máy đại diện Rework thành công; BRAVO có phiếu đề nghị xuất đúng lệnh và đúng mã.',
'3. Đã kiểm tra tồn kho xuất và hoàn tất điều chuyển từ kho chính sang kho phân xưởng.',
'4. Trước khi hoàn thành MES, tồn của đúng mã hàng tại kho phân xưởng đáp ứng số lượng xử lý.',
'5. MES thống kê và hoàn thành thành công; kết quả đã chuyển lên BRAVO.',
'6. BRAVO sinh phiếu xuất chuyển mã giao dịch 2212; kho K, kho J, mã vật tư xuất, danh sách vật tư nhập và TK nợ/TK có đã được đối chiếu.']: p(t)
h('Thông tin cần xác nhận để hoàn thiện quy trình')
p('Đối với 8A và 8B: xác nhận giá trị mục 4 Đối tượng; ánh xạ chính xác các trường b và d khi kiểm tra tồn; nguồn cấu hình kho xuất K và kho nhập J; tài khoản áp dụng; cách liên kết chứng từ với RMA của 8B.')
p('Đối với 7A và 7B: xác nhận điểm tạo lệnh, cách đồng bộ MES, máy đại diện Sorting, quy tắc lấy vật tư và kiểm tra tồn, cách thống kê hoàn thành, loại chứng từ BRAVO phát sinh, mã giao dịch và kết quả đầu ra. Luồng Sorting cần được xác định riêng trước khi áp dụng vận hành; không mặc định sử dụng luồng chuyển mã 2212 của Rework.')
d.core_properties.title='Quy trình xử lý Rework và Sorting trên MES và BRAVO'
d.core_properties.subject='Luồng 8A 8B và phạm vi 7A 7B'
d.core_properties.author=''
d.save(out/'Quy_trinh_Rework_Sorting_MES_BRAVO.docx')
print(out/'Quy_trinh_Rework_Sorting_MES_BRAVO.docx')
