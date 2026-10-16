from pathlib import Path
from copy import deepcopy
import re,hashlib
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt,RGBColor

root=Path(__file__).parent
src=Path('C:/Users/tuanhoang/Downloads/Quy_trinh_Rework_Sorting_MES_BRAVO_Sorting_chi_tiet.docx')
d=Document(src)
def hashes(doc):return [hashlib.sha256(r.target_part.blob).hexdigest() for r in doc.part.rels.values() if 'image' in r.reltype]
before=hashes(d);count=len(d.inline_shapes)
# Pasted screenshots share heading paragraphs in several places. Separate them
# before numbering so screenshots do not appear as blank entries in navigation.
for p in list(d.paragraphs):
    if not p._p.xpath('.//w:drawing|.//w:pict'):continue
    if p.text.strip():
        el=OxmlElement('w:p');p._p.addprevious(el);imgp=Paragraph(el,p._parent)
        for run in list(p._p):
            if run.xpath('.//w:drawing|.//w:pict'):
                copied=deepcopy(run)
                for t in copied.xpath('.//w:t'):t.getparent().remove(t)
                el.append(copied)
                for drawing in run.xpath('.//w:drawing|.//w:pict'):drawing.getparent().remove(drawing)
        imgp.style='Normal';imgp.alignment=1
        imgp.paragraph_format.space_after=Pt(8)
        imgp.paragraph_format.keep_together=True
    else:
        p.style='Normal';p.alignment=1
        p.paragraph_format.keep_with_next=False
        p.paragraph_format.page_break_before=False
        p.paragraph_format.keep_together=True
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(8)
        for r in p.runs:r.font.size=Pt(11)

main=second=third=0;sorting_step=0;figure=0
for p in d.paragraphs:
    text=p.text.strip();style=p.style.name
    if style=='Heading 1':
        main+=1;second=third=0
        text=re.sub(r'^\d+[ .]*','',text)
        p.text=f'{main}. {text}'
    elif style=='Heading 2' and text:
        second+=1;third=0
        text=re.sub(r'^\d+(?:[ .]+\d+)*[ .]*','',text)
        p.text=f'{main}.{second} {text}'
    elif style=='Heading 3' and text:
        third+=1
        text=re.sub(r'^\d+(?:\.\d+)+\s+','',text)
        if main==2 and text.startswith('Bước '):
            sorting_step+=1;text=re.sub(r'^Bước \d+',f'Bước {sorting_step}',text)
        text=text.replace('Màn hình danh sách lệnh RW ở MES XƯỞNG','Màn hình danh sách lệnh Rework trên MES xưởng').replace('Màn hình lấy vật tự','Màn hình lấy vật tư')
        p.text=f'{main}.{second}.{third} {text.rstrip()}'
    elif style=='Caption' and re.match(r'^Hình\s+\d+',text):
        figure+=1;p.text=re.sub(r'^Hình\s+\d+',f'Hình {figure}',text)
    if p.style.name.startswith('Heading'):
        pf=p.paragraph_format;pf.keep_with_next=True;pf.keep_together=True
        pf.space_before=Pt(12 if style!='Heading 3' else 8);pf.space_after=Pt(6)
        for r in p.runs:
            r.font.name='Arial';r.font.color.rgb=RGBColor(0,0,0);r.font.size=None;r.bold=None

for name,size in [('Heading 1',18),('Heading 2',14),('Heading 3',11.5)]:
    st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0)
    # Explicit outline levels make the navigation pane match the displayed hierarchy.
    pr=st.element.get_or_add_pPr()
    for e in list(pr):
        if e.tag==qn('w:outlineLvl'):pr.remove(e)
    e=OxmlElement('w:outlineLvl');e.set(qn('w:val'),str(int(name[-1])-1));pr.append(e)
target=root/'Quy_trinh_Rework_Sorting_MES_BRAVO_Chuan_hoa_tieu_de.docx'
d.save(target)
v=Document(target)
assert count==len(v.inline_shapes)==13
assert before==hashes(v)
assert not any(p.style.name.startswith('Heading') and p._p.xpath('.//w:drawing') for p in v.paragraphs)
print(target)
for p in v.paragraphs:
    if p.style.name.startswith('Heading'):print(p.style.name,p.text)
print('Verified all 13 images preserved')
