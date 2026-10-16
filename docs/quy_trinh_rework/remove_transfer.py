from pathlib import Path
from docx import Document
import re,hashlib
root=Path(__file__).parent
d=Document(root/'Quy_trinh_Rework_Sorting_MES_BRAVO_Chuan_hoa_tieu_de.docx')
def hashes(doc):return [hashlib.sha256(r.target_part.blob).hexdigest() for r in doc.part.rels.values() if 'image' in r.reltype]
before=hashes(d)
start=next(p for p in d.paragraphs if p.text=='1.4 Điều chuyển vật tư sang kho phân xưởng')
node=start._p
while node is not None:
    if node is not start._p and node.tag.endswith('}p'):
        text=''.join(node.itertext())
        if '1.5 Thống kê và hoàn thành lệnh trên MES' in text:break
    nxt=node.getnext();node.getparent().remove(node);node=nxt
in_rework=False
for p in d.paragraphs:
    if p.style.name=='Heading 1':in_rework=p.text.startswith('1.')
    if not in_rework:continue
    text=p.text
    if p.style.name.startswith('Heading'):
        text=re.sub(r'^1\.([56])(?=[ .])',lambda m:'1.'+str(int(m.group(1))-1),text)
        text=re.sub(r'Bước ([567])',lambda m:'Bước '+str(int(m.group(1))-1),text)
        if text!=p.text:p.text=text
    # Remove the matching deleted activity from the overview and close numbering.
    elif text=='5. Trên BRAVO, sử dụng đề nghị xuất để điều chuyển từ kho chính sang kho phân xưởng.':
        p._p.getparent().remove(p._p)
    elif re.match(r'^[67]\. ',text):
        p.text=str(int(text[0])-1)+text[1:]
target=root/'Quy_trinh_Rework_Sorting_MES_BRAVO_Da_bo_muc_dieu_chuyen.docx'
d.save(target)
v=Document(target)
assert before==hashes(v) and len(v.inline_shapes)==13
assert not any('Điều chuyển' in p.text for p in v.paragraphs if p.style.name.startswith('Heading'))
print(target)
for p in v.paragraphs:
    if p.style.name=='Heading 2':print(p.text)
