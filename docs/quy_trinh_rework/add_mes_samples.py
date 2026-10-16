from pathlib import Path
from docx import Document
from docx.shared import Pt
import hashlib

root=Path(__file__).parent
src=root/'Quy_trinh_Rework_Sorting_MES_BRAVO_Can_chinh.docx'
d=Document(src)
def image_hashes(doc):
    return [hashlib.sha256(r.target_part.blob).hexdigest() for r in doc.part.rels.values() if 'image' in r.reltype]
before=image_hashes(d)
p=d.add_heading('3 Mô tả giao diện mẫu mong muốn xây dựng trên MES',1)
p.paragraph_format.page_break_before=True
d.add_paragraph('Phần này trình bày ảnh mẫu và yêu cầu giao diện MES theo hai nhóm Rework và Sorting. Mỗi nhóm ghi rõ màn hình cần xây dựng, thông tin cần hiển thị, thao tác mong muốn và kết quả xử lý.')
for number,name,scope in [('3 1','Rework','8A Rework nội bộ và 8B Rework hàng bán bị trả lại theo RMA'),('3 2','Sorting','7A Sorting nội bộ và 7B Sorting hàng bán bị trả lại theo RMA')]:
    p=d.add_heading(f'{number} {name}',2)
    p.paragraph_format.page_break_before=(number=='3 2')
    d.add_paragraph(f'Áp dụng cho {scope}.')
    p=d.add_paragraph('Ảnh mẫu giao diện MES mong muốn','Caption')
    p.paragraph_format.keep_with_next=True
    p=d.add_paragraph(f'[Chèn ảnh mẫu giao diện MES cho nhóm {name} tại đây]')
    p.alignment=1
    p.paragraph_format.space_before=Pt(12)
    p.paragraph_format.space_after=Pt(100)
    p.runs[0].italic=True
    for heading,placeholder in [
        ('Màn hình và mục đích sử dụng','[Bổ sung tên màn hình, người sử dụng và mục đích của màn hình]'),
        ('Thông tin cần hiển thị','[Bổ sung các trường dữ liệu, danh sách, giá trị mặc định và thông tin cần đẩy đúng]'),
        ('Thao tác và logic xử lý mong muốn','[Bổ sung thao tác người dùng, điều kiện kiểm tra, phản hồi và kết quả mong muốn trên MES hoặc BRAVO]'),
        ('Chú ý khi xây dựng','[Bổ sung điểm cần lưu ý và yêu cầu riêng theo ảnh mẫu]')]:
        d.add_heading(heading,3)
        d.add_paragraph(placeholder)
target=root/'Quy_trinh_Rework_Sorting_MES_BRAVO_Bo_sung_muc_3.docx'
d.save(target)
v=Document(target)
assert before==image_hashes(v)
assert len(v.inline_shapes)==4
assert len([p for p in v.paragraphs if p.style.name=='Heading 1'])==3
print(target)
print('Verified 3 main sections and all 4 existing images preserved')
